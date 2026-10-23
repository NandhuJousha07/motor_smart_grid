import openpyxl,re,sys,json
SHEETS={'TN PO':'TN','KL PO':'KL','KA PO':'KA','AP-TS PO':'AP-TS','MH PO':'MH'}
def norm(v):
    if v is None: return ''
    s=str(v)
    if isinstance(v,float) and v.is_integer(): s=str(int(v))
    return re.sub(r'\s+',' ',s).strip()
def classify_header(cells):
    m={}
    for i,c in enumerate(cells):
        t=norm(c).lower()
        if not t: continue
        if t=='product': m['product']=i
        elif 'categ' in t: m['cat']=i
        elif t.startswith('company'): m['company']=i
        elif t.startswith('age'): m['age']=i
        elif t.startswith('insp'): m['insp']=i
        elif t.startswith('discount'): m['disc']=i
        elif t.startswith('pack'): m['pkg']=i
        elif t=='tp': m['tp']=i
        elif t.startswith('remark'): m['remark']=i
        elif t.startswith('rto'): m['rto']=i
    return m
def parse(path):
    wb=openpyxl.load_workbook(path,data_only=True)
    rows=[]
    for ws in wb.worksheets:
        if ws.title not in SHEETS or ws.sheet_state!='visible': continue
        st=SHEETS[ws.title]; hdr=None; block=None
        for i,r in enumerate(ws.iter_rows(values_only=True),1):
            vals=[v for v in r if norm(v)]
            if not vals: continue
            h=classify_header(r)
            if 'company' in h and 'cat' in h and ('pkg' in h or 'tp' in h):
                hdr=h; continue
            if len(vals)<=2:
                t=norm(vals[-1]).upper()
                if 'CAR' in t: block='PC'
                elif t in('TW','TWO WHEELER'): block='TW'
                elif 'CV' in t or 'COMMERCIAL' in t: block='CV'
                continue
            if hdr is None: continue
            g=lambda k: r[hdr[k]] if k in hdr and hdr[k]<len(r) else None
            prod=norm(g('product'))
            if 'product' not in hdr:
                # product may be in the column before category
                c=hdr['cat']
                for j in range(c-1,-1,-1):
                    if norm(r[j]): prod=norm(r[j]); break
            rows.append(dict(state=st,sheet=ws.title,row=i,block=block,product=prod,cat=norm(g('cat')),company=norm(g('company')),
               age=norm(g('age')),insp=norm(g('insp')),disc=norm(g('disc')),disc_raw=g('disc'),pkg=g('pkg'),tp=g('tp'),remark=norm(g('remark')),rto=norm(g('rto'))))
    return rows
STATE_NAMES={'TAMIL NADU':'TN','TAMILNADU':'TN','TN':'TN','KERALA':'KL','KL':'KL','KARNATAKA':'KA','KA':'KA',
  'ANDHRA PRADESH':'AP-TS','TELANGANA':'AP-TS','ANDHRA PRADESH & TELANGANA':'AP-TS','AP-TS':'AP-TS','AP':'AP-TS','TS':'AP-TS',
  'MAHARASHTRA':'MH','MH':'MH'}
def parse_supplement(path):
    """Extra GRID rows sent separately (e.g. an insurer's missing rates).
    Same columns as the state sheets plus a 'State' column; the column before
    'Vehicle Category' says Pvt Car / TW / CV."""
    wb=openpyxl.load_workbook(path,data_only=True)
    rows=[]
    for ws in wb.worksheets:
        if ws.sheet_state!='visible': continue
        hdr=None; sc=None
        for i,r in enumerate(ws.iter_rows(values_only=True),1):
            if not any(norm(v) for v in r): continue
            h=classify_header(r)
            if 'company' in h and 'cat' in h and ('pkg' in h or 'tp' in h):
                hdr=h; sc=next((j for j,c in enumerate(r) if norm(c).lower()=='state'),None); continue
            if hdr is None: continue
            g=lambda k: r[hdr[k]] if k in hdr and hdr[k]<len(r) else None
            st=STATE_NAMES.get(norm(r[sc]).upper()) if sc is not None else None
            if not st: raise SystemExit('%s row %d: unknown state %r' % (ws.title,i,r[sc] if sc is not None else None))
            prod=''
            for j in range(hdr['cat']-1,-1,-1):
                if j!=sc and norm(r[j]): prod=norm(r[j]); break
            P=prod.upper()
            block='PC' if ('CAR' in P or 'PVT' in P) else ('TW' if P in ('TW','TWO WHEELER') else 'CV')
            rows.append(dict(state=st,sheet=ws.title,row=i,block=block,product=prod,cat=norm(g('cat')),company=norm(g('company')),
               age=norm(g('age')),insp=norm(g('insp')),disc=norm(g('disc')),disc_raw=g('disc'),pkg=g('pkg'),tp=g('tp'),remark=norm(g('remark')),rto=norm(g('rto')),
               supplement=True))
    return rows
if __name__=='__main__':
    rows=parse(sys.argv[1]); json.dump(rows,open(sys.argv[2],'w'),default=str)
    import collections
    print(len(rows),collections.Counter((r['state'],r['block'],r['product']) for r in rows))
